"""Isolated phase-one rotation decisions. Never writes catalog/pricing tables."""
from __future__ import annotations

import fcntl
import hashlib
import io
import json
import math
import secrets
import sqlite3
from collections import defaultdict
from datetime import date, datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Protection

from planning_center.db import get_connection, normalize_season_year, utc_now

SHANGHAI = ZoneInfo("Asia/Shanghai")
TMALL = {"天猫", "天猫官", "天猫奥"}
CHANNELS = TMALL | {"唯品", "同款"}
STATE_LABELS = {"draft": "草案", "confirmed": "已确认待执行", "executing": "执行中", "completed": "已完成（人工反馈）", "cancelled": "已取消"}
SALES_HEADERS = ("款号", "销售体系", "开始日期", "结束日期", "净销量")
SCHEMA = (
    "CREATE TABLE IF NOT EXISTS rotation_meta (key TEXT PRIMARY KEY,value TEXT NOT NULL)",
    "CREATE TABLE IF NOT EXISTS rotation_sources (id INTEGER PRIMARY KEY, payload TEXT NOT NULL, synced_at TEXT NOT NULL)",
    "CREATE TABLE IF NOT EXISTS rotation_syncs (id INTEGER PRIMARY KEY, synced_at TEXT NOT NULL, item_count INTEGER NOT NULL, actor TEXT NOT NULL)",
    "CREATE TABLE IF NOT EXISTS rotation_sales_batches (id INTEGER PRIMARY KEY, actor TEXT NOT NULL, created_at TEXT NOT NULL, row_count INTEGER NOT NULL)",
    "CREATE TABLE IF NOT EXISTS rotation_sales (id INTEGER PRIMARY KEY, batch_id INTEGER NOT NULL REFERENCES rotation_sales_batches(id), style_code TEXT NOT NULL, source_ids TEXT NOT NULL, platform TEXT NOT NULL, start_date TEXT NOT NULL, end_date TEXT NOT NULL, quantity INTEGER NOT NULL, active INTEGER NOT NULL DEFAULT 1)",
    "CREATE INDEX IF NOT EXISTS idx_rotation_sales_lookup ON rotation_sales(style_code, platform, active, start_date, end_date)",
    "CREATE TABLE IF NOT EXISTS rotation_plans (id INTEGER PRIMARY KEY, state TEXT NOT NULL DEFAULT 'draft', criteria TEXT NOT NULL, actor TEXT NOT NULL, created_at TEXT NOT NULL, confirmed_by TEXT, confirmed_at TEXT, updated_at TEXT NOT NULL)",
    "CREATE TABLE IF NOT EXISTS rotation_plan_items (id INTEGER PRIMARY KEY, plan_id INTEGER NOT NULL REFERENCES rotation_plans(id), style_code TEXT NOT NULL, snapshot TEXT NOT NULL, target TEXT NOT NULL DEFAULT '', active INTEGER NOT NULL DEFAULT 1, executed_at TEXT, feedback TEXT NOT NULL DEFAULT '', UNIQUE(plan_id, style_code))",
    "CREATE INDEX IF NOT EXISTS idx_rotation_items_style ON rotation_plan_items(style_code, active, plan_id)",
    "CREATE TABLE IF NOT EXISTS rotation_item_sources (item_id INTEGER NOT NULL REFERENCES rotation_plan_items(id), source_id INTEGER NOT NULL, PRIMARY KEY(item_id,source_id))",
    "CREATE INDEX IF NOT EXISTS idx_rotation_item_sources_id ON rotation_item_sources(source_id, item_id)",
    "CREATE TABLE IF NOT EXISTS rotation_events (id INTEGER PRIMARY KEY, plan_id INTEGER NOT NULL, actor TEXT NOT NULL, action TEXT NOT NULL, details TEXT NOT NULL, created_at TEXT NOT NULL)",
)


def ensure_schema(db_path):
    """Back up an existing DB before the first additive migration; never prune."""
    path = Path(db_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.with_suffix(path.suffix + ".rotation-migration.lock").open("a") as lock:
        fcntl.flock(lock, fcntl.LOCK_EX)
        with get_connection(path) as connection:
            if connection.execute("SELECT 1 FROM sqlite_master WHERE name='rotation_sources'").fetchone():
                return
            backup_info = {"required": False, "verified": True}
            if connection.execute("SELECT 1 FROM sqlite_master WHERE type='table' LIMIT 1").fetchone():
                folder = path.parent / "backups"
                folder.mkdir(mode=0o700, exist_ok=True)
                target = folder / ("pre-rotation-" + datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S") + "-" + secrets.token_hex(4) + ".db")
                target.touch(mode=0o600, exist_ok=False)
                with sqlite3.connect(target) as backup:
                    connection.backup(backup)
                    if backup.execute("PRAGMA quick_check").fetchone()[0] != "ok":
                        raise RuntimeError("升级前数据库备份校验失败，已停止升级。")
                backup_info = {"required": True, "verified": True, "path": str(target), "created_at": utc_now()}
            with connection:
                connection.execute("BEGIN IMMEDIATE")
                for sql in SCHEMA:
                    connection.execute(sql)
                connection.execute("INSERT INTO rotation_meta(key,value) VALUES('migration_backup',?)", (json.dumps(backup_info),))


def readiness(db_path):
    with get_connection(db_path) as connection:
        names = {r[0] for r in connection.execute("SELECT name FROM sqlite_master WHERE type='table' AND name LIKE 'rotation_%'")}
        ready = {"rotation_meta", "rotation_sources", "rotation_syncs", "rotation_sales", "rotation_sales_batches", "rotation_plans", "rotation_plan_items", "rotation_item_sources", "rotation_events"}.issubset(names)
        row = connection.execute("SELECT value FROM rotation_meta WHERE key='migration_backup'").fetchone() if ready else None
        backup = json.loads(row[0]) if row else {}
        verified = bool(backup.get("verified")) and (not backup.get("required") or Path(backup.get("path", "")).is_file())
        return {"schema_ready": ready, "migration_backup_verified": verified, "migration_backup_required": bool(backup.get("required"))}


def timestamp(value):
    value = str(value or "").strip()
    if not value:
        return None
    parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed.astimezone(timezone.utc)


def local_day(value):
    stamp = timestamp(value)
    return stamp.astimezone(SHANGHAI).date() if stamp else None


def today():
    return datetime.now(SHANGHAI).date()


def fingerprint(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()).hexdigest()


def catalog_page(db_path, after=0, limit=500):
    from catalog_backend.rotation import catalog_page as read_page
    return read_page(db_path, after, limit)


def sync_sources(db_path, items, actor):
    try:
        ids = [int(item["id"]) for item in items]
        for item in items:
            if int(item["id"]) <= 0 or not isinstance(item.get("receipts"), list):
                raise ValueError()
            for key in ("style_code", "style_color", "launch_channel", "status", "lifecycle_status", "product_line"):
                if not isinstance(item.get(key), str):
                    raise ValueError()
            int(item.get("c_release_no") or 0)
            timestamp(item.get("recall_at"))
            for receipt in item["receipts"]:
                int(receipt["release_no"])
                if not timestamp(receipt["received_at"]):
                    raise ValueError()
    except (KeyError, TypeError, ValueError) as error:
        raise ValueError("轮转资料接口返回无效字段，原有资料未改动。") from error
    if len(ids) != len(set(ids)):
        raise ValueError("轮转资料接口返回重复编号，未更新本地资料。")
    now = utc_now()
    with get_connection(db_path) as connection:
        connection.execute("BEGIN IMMEDIATE")
        # Keep withdrawn/missing rows as evidence, but exclude them from use.
        connection.execute("UPDATE rotation_sources SET synced_at = ''")
        connection.executemany(
            "INSERT INTO rotation_sources(id,payload,synced_at) VALUES(?,?,?) ON CONFLICT(id) DO UPDATE SET payload=excluded.payload,synced_at=excluded.synced_at",
            [(item["id"], json.dumps(item, ensure_ascii=False), now) for item in items],
        )
        connection.execute("INSERT INTO rotation_syncs(synced_at,item_count,actor) VALUES(?,?,?)", (now, len(items), actor))
    return len(items)


def grouped_styles(connection):
    grouped = defaultdict(list)
    for row in connection.execute("SELECT payload FROM rotation_sources WHERE synced_at != ''"):
        item = json.loads(row[0])
        if item.get("product_line") == "red" and str(item.get("style_code") or "").strip():
            grouped[str(item["style_code"]).strip()].append(item)
    result = []
    for code, colors in grouped.items():
        # A recall on any color resets the clock for the ENTIRE style. Do not
        # let an untouched color's old receipt restore a pre-recall date.
        colors.sort(key=lambda c: c["id"])
        recall = max((str(c.get("recall_at") or "") for c in colors if c.get("recall_at")), default="", key=timestamp)
        eligible_receipts = []
        channels = {str(c.get("launch_channel") or "") for c in colors}
        blocked = any(c.get("status") not in {"published", "received"} or c.get("lifecycle_status") != "active" or c.get("workflow_restart_required") for c in colors)
        for color in colors:
            for receipt in color.get("receipts", []):
                if receipt.get("recipient_platform") not in {"tmall", "vip"}:
                    continue
                if int(receipt.get("release_no") or 0) != int(color.get("c_release_no") or 0):
                    continue
                stamp = str(receipt.get("received_at") or "")
                if stamp and (not recall or timestamp(stamp) >= timestamp(recall)):
                    eligible_receipts.append(stamp)
        reference = min(eligible_receipts, default="", key=timestamp)
        reason = ""
        if blocked:
            reason = "整款含召回、未提交运营或停用资料"
        elif len(channels) != 1 or not channels.issubset(CHANNELS):
            reason = "同款渠道不一致或不受支持，请先核对"
        elif not reference:
            reason = "召回更新后待接收" if recall else "尚无天猫／唯品有效接收记录"
        result.append({
            "style_code": code, "source_ids": sorted(int(c["id"]) for c in colors),
            "colors": [str(c.get("style_color") or "") for c in colors],
            "product_names": sorted({str(c.get("product_name") or "") for c in colors}),
            "season_years": sorted({normalize_season_year(c.get("season_year")) for c in colors}),
            "suppliers": sorted({str(c.get("supplier") or "") for c in colors}),
            "categories": sorted({str(c.get("category") or "") for c in colors}),
            "channel": next(iter(channels)) if len(channels) == 1 else "不一致",
            "reference_at": reference, "recall_at": recall, "reason": reason,
            "receipt_evidence": [{"source_product_id": c["id"], "release_no": c.get("c_release_no"), "receipts": c.get("receipts", []), "recall_at": c.get("recall_at", "")} for c in colors],
            # Cosmetic metadata updates must not force a new decision. Stable
            # IDs, channel, release/recall and receipt evidence remain binding.
            "source_signature": fingerprint([{k: c.get(k) for k in ("id", "style_code", "style_color", "launch_channel", "status", "lifecycle_status", "c_release_no", "workflow_restart_required", "recall_at", "receipts")} for c in colors]),
        })
    return sorted(result, key=lambda s: s["style_code"])


def clean_criteria(values):
    mode = str(values.get("mode") or "time")
    if mode not in {"time", "sales", "both"}:
        raise ValueError("筛选模式不正确。")
    age_value = values.get("min_days")
    age = int(70 if age_value in (None, "") else age_value)
    if not 0 <= age <= 10000:
        raise ValueError("接收天数须为 0—10000。")
    result = {"mode": mode, "min_days": age, "season_year": str(values.get("season_year") or "").strip(), "channel": str(values.get("channel") or "").strip()}
    for key in ("supplier", "category"):
        raw = values.get(key) or []
        result[key] = list(dict.fromkeys(raw if isinstance(raw, (list, tuple)) else [raw]))
    if mode in {"sales", "both"}:
        start, end = date.fromisoformat(str(values.get("start_date") or "")), date.fromisoformat(str(values.get("end_date") or ""))
        if start > end or end > today():
            raise ValueError("销售期间须已结束，且开始日期不能晚于结束日期。")
        maximum = int(values.get("max_sales") or 0)
        result.update(start_date=start.isoformat(), end_date=end.isoformat(), max_sales=maximum)
    return result


def candidates(db_path, criteria, evaluation_day=None):
    criteria = clean_criteria(criteria)
    evaluation_day = evaluation_day or today()
    with get_connection(db_path) as connection:
        connection.execute("BEGIN")
        styles = grouped_styles(connection)
        sales = defaultdict(dict)
        if criteria["mode"] in {"sales", "both"}:
            for row in connection.execute("SELECT * FROM rotation_sales WHERE active=1 AND start_date=? AND end_date=?", (criteria["start_date"], criteria["end_date"])):
                sales[row["style_code"]][row["platform"]] = dict(row)
        busy = {row[0] for row in connection.execute("SELECT i.style_code FROM rotation_plan_items i JOIN rotation_plans p ON p.id=i.plan_id WHERE i.active=1 AND p.state IN ('draft','confirmed','executing')")}
        busy_sources = {row[0] for row in connection.execute("SELECT s.source_id FROM rotation_item_sources s JOIN rotation_plan_items i ON i.id=s.item_id JOIN rotation_plans p ON p.id=i.plan_id WHERE i.active=1 AND p.state IN ('draft','confirmed','executing')")}
        done = defaultdict(list)
        for row in connection.execute("SELECT snapshot FROM rotation_plan_items WHERE executed_at IS NOT NULL"):
            old = json.loads(row[0])
            for sid in old["source_ids"]:
                done[sid].append(old)
        for style in styles:
            style["age_days"] = (evaluation_day - local_day(style["reference_at"])).days if style["reference_at"] else None
            style["sales"] = {p: sale for p, sale in sales.get(style["style_code"], {}).items() if json.loads(sale["source_ids"]) == style["source_ids"]}
            required = ["tmall", "vip"] if style["channel"] == "同款" else ["vip"] if style["channel"] == "唯品" else ["tmall"]
            style["sales_quantity"] = sum(style["sales"][p]["quantity"] for p in required) if all(p in style["sales"] for p in required) else None
            style["busy"] = style["style_code"] in busy or bool(busy_sources.intersection(style["source_ids"]))
            old_rounds = [old for sid in style["source_ids"] for old in done[sid]]
            for old in old_rounds:
                # Phase one does not pretend that manual feedback has changed
                # platform listings. Suppress the same source round until new
                # catalog channel/recall/release evidence exists.
                before = {e["source_product_id"]: (e["release_no"], e["recall_at"]) for e in old["receipt_evidence"]}
                current = {e["source_product_id"]: (e["release_no"], e["recall_at"]) for e in style["receipt_evidence"]}
                if old["channel"] == style["channel"] and all(current.get(sid) == evidence for sid, evidence in before.items()):
                    style["reason"] = "本轮已反馈执行，暂不重复推荐；不代表平台已实际换渠道"
                    break
        result = []
        for style in styles:
            if style["reason"] or style["busy"]:
                continue
            if criteria["season_year"] and criteria["season_year"] not in style["season_years"]:
                continue
            if criteria["channel"] and criteria["channel"] != style["channel"]:
                continue
            if any(criteria[k] and not set(criteria[k]).intersection(style[out]) for k, out in (("supplier", "suppliers"), ("category", "categories"))):
                continue
            if criteria["mode"] in {"time", "both"} and (style["age_days"] is None or style["age_days"] < criteria["min_days"]):
                continue
            if criteria["mode"] in {"sales", "both"} and (style["sales_quantity"] is None or style["sales_quantity"] > criteria["max_sales"]):
                continue
            result.append(style)
        return result


def valid_target(channel, target):
    return (channel in TMALL and target == "唯品") or (channel == "唯品" and target in TMALL) or (channel == "同款" and target in (TMALL | {"唯品"}))


def add_event(connection, plan_id, actor, action, details=""):
    connection.execute("INSERT INTO rotation_events(plan_id,actor,action,details,created_at) VALUES(?,?,?,?,?)", (plan_id, actor, action, details, utc_now()))


def create_plan(db_path, criteria, selected, target, actor, expected_signature):
    rows = candidates(db_path, criteria)
    if expected_signature != fingerprint(rows):
        raise ValueError("筛选资料已变化，请重新预览后生成方案。")
    if selected is not None:
        selected = set(selected)
        rows = [r for r in rows if r["style_code"] in selected]
        if len(rows) != len(selected):
            raise ValueError("所选款号已不满足条件，请重新预览。")
    if not rows:
        raise ValueError("没有可生成方案的资料。")
    if len(rows) > 10000:
        raise ValueError("单个方案最多 10000 款，请缩小筛选范围。")
    for row in rows:
        if not valid_target(row["channel"], target):
            raise ValueError(f"{row['style_code']}：目标渠道不符合跨体系／同款转单渠道规则，请按当前渠道分批生成。")
    with get_connection(db_path) as connection:
        connection.execute("BEGIN IMMEDIATE")
        live_sources = {s["style_code"]: s for s in grouped_styles(connection)}
        busy_ids = {r[0] for r in connection.execute("SELECT s.source_id FROM rotation_item_sources s JOIN rotation_plan_items i ON i.id=s.item_id JOIN rotation_plans p ON p.id=i.plan_id WHERE i.active=1 AND p.state IN ('draft','confirmed','executing')")}
        for row in rows:
            live = connection.execute("SELECT 1 FROM rotation_plan_items i JOIN rotation_plans p ON p.id=i.plan_id WHERE i.style_code=? AND i.active=1 AND p.state IN ('draft','confirmed','executing')", (row["style_code"],)).fetchone()
            source = live_sources.get(row["style_code"])
            if live or busy_ids.intersection(row["source_ids"]) or not source or source["source_signature"] != row["source_signature"]:
                raise ValueError("存在进行中的方案或资料已变化，请重新预览。")
            for sale in row["sales"].values():
                if not connection.execute("SELECT 1 FROM rotation_sales WHERE id=? AND active=1", (sale["id"],)).fetchone():
                    raise ValueError("销售资料已变化，请重新预览。")
        now = utc_now()
        plan_id = connection.execute("INSERT INTO rotation_plans(criteria,actor,created_at,updated_at) VALUES(?,?,?,?)", (json.dumps(clean_criteria(criteria), ensure_ascii=False), actor, now, now)).lastrowid
        connection.executemany("INSERT INTO rotation_plan_items(plan_id,style_code,snapshot,target) VALUES(?,?,?,?)", [(plan_id, row["style_code"], json.dumps(row, ensure_ascii=False), target) for row in rows])
        item_ids = {r["style_code"]: r["id"] for r in connection.execute("SELECT id,style_code FROM rotation_plan_items WHERE plan_id=?", (plan_id,))}
        connection.executemany("INSERT INTO rotation_item_sources(item_id,source_id) VALUES(?,?)", [(item_ids[r["style_code"]], sid) for r in rows for sid in r["source_ids"]])
        add_event(connection, plan_id, actor, "生成草案", f"整款 {len(rows)} 款")
    return plan_id


def plan_detail(db_path, plan_id):
    with get_connection(db_path) as connection:
        connection.execute("BEGIN")
        plan = connection.execute("SELECT * FROM rotation_plans WHERE id=?", (plan_id,)).fetchone()
        if not plan:
            raise LookupError("轮转方案不存在。")
        result = dict(plan)
        result["items"] = [dict(r) for r in connection.execute("SELECT * FROM rotation_plan_items WHERE plan_id=? ORDER BY id", (plan_id,))]
        result["events"] = [dict(r) for r in connection.execute("SELECT * FROM rotation_events WHERE plan_id=? ORDER BY id", (plan_id,))]
        for item in result["items"]:
            item["snapshot"] = json.loads(item["snapshot"])
        return result


def plan_signature(plan):
    return fingerprint([plan["state"], [(i["id"], i["target"], i["active"], i["executed_at"]) for i in plan["items"]]])


def update_plan(db_path, plan_id, action, user, values):
    if user.get("role") not in {"planner", "admin"}:
        raise PermissionError("只有商品部人员可以处理轮转方案。")
    actor = f"{user['id']}:{user['display_name']}"
    with get_connection(db_path) as connection:
        connection.execute("BEGIN IMMEDIATE")
        plan = connection.execute("SELECT * FROM rotation_plans WHERE id=?", (plan_id,)).fetchone()
        if not plan:
            raise LookupError("轮转方案不存在。")
        items = [dict(r) for r in connection.execute("SELECT * FROM rotation_plan_items WHERE plan_id=? AND active=1", (plan_id,))]
        state = plan["state"]
        if values.get("plan_signature"):
            all_items = [dict(r) for r in connection.execute("SELECT * FROM rotation_plan_items WHERE plan_id=? ORDER BY id", (plan_id,))]
            if values["plan_signature"] != plan_signature({"state": state, "items": all_items}):
                raise ValueError("方案已被其他同事更新，请刷新后再操作。")
        if action == "save" and state == "draft":
            for item in items:
                target = str(values.get(f"target_{item['id']}") or item["target"])
                if not valid_target(json.loads(item["snapshot"])["channel"], target):
                    raise ValueError("目标渠道不符合轮转规则。")
                connection.execute("UPDATE rotation_plan_items SET target=?, active=? WHERE id=?", (target, 0 if values.get(f"remove_{item['id']}") == "1" else 1, item["id"]))
        elif action == "confirm" and state == "draft":
            if user.get("role") != "admin":
                raise PermissionError("轮转方案由企划管理员确认。")
            if not items:
                raise ValueError("空方案不能确认。")
            for item in items:
                if (values.get(f"target_{item['id']}") not in (None, item["target"]) or values.get(f"remove_{item['id']}") == "1"):
                    raise ValueError("草案还有未保存的调整，请先保存再确认。")
            current = {s["style_code"]: s for s in grouped_styles(connection)}
            for item in items:
                old = json.loads(item["snapshot"])
                live = current.get(item["style_code"])
                if not live or live["reason"] or live["source_signature"] != old["source_signature"]:
                    raise ValueError(f"{item['style_code']} 来源资料已变化，请取消草案后重新生成。")
                for sale in old.get("sales", {}).values():
                    if not connection.execute("SELECT 1 FROM rotation_sales WHERE id=? AND active=1", (sale["id"],)).fetchone():
                        raise ValueError("销售资料已被新批次替换，请重新生成草案。")
            state = "confirmed"
            connection.execute("UPDATE rotation_plans SET confirmed_by=?,confirmed_at=? WHERE id=?", (actor, utc_now(), plan_id))
        elif action == "start" and state == "confirmed":
            state = "executing"
        elif action == "feedback" and state in {"confirmed", "executing"}:
            selected = [i for i in items if values.get(f"done_{i['id']}") == "1"]
            if not selected:
                raise ValueError("请勾选已人工执行的整款。")
            feedback = str(values.get("feedback") or "").strip()
            if not feedback or len(feedback) > 1000:
                raise ValueError("请填写不超过 1000 字的人工执行说明。")
            for item in selected:
                if not item["executed_at"]:
                    connection.execute("UPDATE rotation_plan_items SET executed_at=?,feedback=? WHERE id=?", (utc_now(), feedback, item["id"]))
            remaining = connection.execute("SELECT COUNT(*) FROM rotation_plan_items WHERE plan_id=? AND active=1 AND executed_at IS NULL", (plan_id,)).fetchone()[0]
            state = "executing" if remaining else "completed"
        elif action == "cancel" and state in {"draft", "confirmed", "executing"}:
            state = "cancelled"
        else:
            raise ValueError("当前方案状态不允许此操作。")
        connection.execute("UPDATE rotation_plans SET state=?,updated_at=? WHERE id=?", (state, utc_now(), plan_id))
        add_event(connection, plan_id, actor, action, json.dumps({k: v for k, v in values.items() if k != "csrf"}, ensure_ascii=False))


def workbook_bytes(headers, rows, title):
    book = Workbook()
    sheet = book.active
    sheet.title = title
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
    for row in sheet:
        for cell in row:
            cell.protection = Protection(locked=False)
            if isinstance(cell.value, str):
                cell.data_type = "s"  # Prevent formula injection from source text.
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="315447")
    sheet.protection.sheet = False
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 22
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def import_sales(db_path, content, actor, replace=False):
    if len(content) > 10 * 1024 * 1024:
        raise ValueError("销量 Excel 不能超过 10MB。")
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError("Excel 解压后过大。")
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        sheet = book.active
        iterator = sheet.iter_rows(values_only=True)
        if tuple(next(iterator, ())) != SALES_HEADERS:
            raise ValueError("请使用轮转销量模板，保留表头顺序。")
        rows, keys, intervals = [], set(), defaultdict(list)
        for number, values in enumerate(iterator, 2):
            if not any(v is not None for v in values):
                continue
            if len(rows) >= 10000 or len(values) != 5:
                raise ValueError("单次最多导入 10000 行，且必须为模板的五列。")
            code, platform, start, end, quantity = values
            code = str(code or "").strip()
            platform = {"天猫体系": "tmall", "唯品体系": "vip"}.get(str(platform or "").strip())
            if not code or not platform:
                raise ValueError(f"第 {number} 行：款号或销售体系无效。")
            start = start.date() if isinstance(start, datetime) else date.fromisoformat(str(start))
            end = end.date() if isinstance(end, datetime) else date.fromisoformat(str(end))
            if start > end or end > today():
                raise ValueError(f"第 {number} 行：销售期间无效。")
            if isinstance(quantity, bool) or not isinstance(quantity, (int, float)) or not math.isfinite(quantity) or quantity != int(quantity):
                raise ValueError(f"第 {number} 行：净销量必须是整数，缺失不能当作 0。")
            key = (code, platform, start.isoformat(), end.isoformat())
            if key in keys:
                raise ValueError(f"第 {number} 行：重复款号、体系和销售期间。")
            keys.add(key)
            rows.append((*key, int(quantity)))
            intervals[(code, platform)].append((key[2], key[3]))
        book.close()
        for periods in intervals.values():
            periods.sort()
            if any(right[0] <= left[1] for left, right in zip(periods, periods[1:])):
                raise ValueError("同批次的同款同体系销售期间不能重叠。")
    except (ValueError, TypeError) as error:
        raise ValueError(f"销量导入失败：{error}") from error
    except Exception as error:
        raise ValueError("无法读取销量 Excel，请使用有效的 .xlsx 模板。") from error
    if not rows:
        raise ValueError("销量表没有数据。")
    with get_connection(db_path) as connection:
        connection.execute("BEGIN IMMEDIATE")
        known = {s["style_code"]: s["source_ids"] for s in grouped_styles(connection)}
        for code, platform, start, end, quantity in rows:
            if code not in known:
                raise ValueError(f"款号 {code} 尚未同步，请先同步并核对款号。")
            overlaps = connection.execute("SELECT * FROM rotation_sales WHERE style_code=? AND platform=? AND active=1 AND start_date<=? AND end_date>=?", (code, platform, end, start)).fetchall()
            if overlaps and (not replace or any(r["start_date"] != start or r["end_date"] != end for r in overlaps)):
                raise ValueError(f"{code} 销售期间已存在或重叠。仅允许显式替换完全相同期间，历史会保留。")
            for old in overlaps:
                connection.execute("UPDATE rotation_sales SET active=0 WHERE id=?", (old["id"],))
        batch_id = connection.execute("INSERT INTO rotation_sales_batches(actor,created_at,row_count) VALUES(?,?,?)", (actor, utc_now(), len(rows))).lastrowid
        connection.executemany("INSERT INTO rotation_sales(batch_id,style_code,source_ids,platform,start_date,end_date,quantity) VALUES(?,?,?,?,?,?,?)", [(batch_id, r[0], json.dumps(known[r[0]]), *r[1:]) for r in rows])
    return len(rows)
