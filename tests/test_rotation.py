from __future__ import annotations

import io
import json
import sqlite3
import tempfile
import time
import unittest
from concurrent.futures import ThreadPoolExecutor
from contextlib import closing
from datetime import date
from pathlib import Path
from unittest.mock import patch
from urllib.parse import urlencode

from openpyxl import load_workbook

from catalog_backend import CatalogApplication, db as catalog_db
from planning_center import PlanningApplication, db
from planning_center import rotation as rot
from tests.test_planning import PlanningCenterTests as _PlanningCenterTests


class RotationTests(unittest.TestCase):
    wsgi_request = _PlanningCenterTests.wsgi_request
    login_cookie = _PlanningCenterTests.login_cookie
    workbook_multipart = _PlanningCenterTests.workbook_multipart

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.path = self.root / "planning.db"
        db.init_db(self.path)
        self.app = PlanningApplication(self.path, "http://catalog.invalid", "test-token")
        self.planner = db.authenticate_user(self.path, "planner", "demo123")
        self.admin = db.authenticate_user(self.path, "planning_admin", "demo123")

    def tearDown(self):
        self.temp.cleanup()

    def source(self, pid=1, code="A001", channel="天猫官", stamp="2026-01-01T01:00:00Z", **updates):
        item = dict(id=pid, style_code=code, style_color=f"{code}-{pid}", product_name="测试商品", season_year="2026秋冬", supplier="供应商A", category="针织衫", launch_channel=channel, status="received", lifecycle_status="active", product_line="red", c_release_no=1, current_version_no=2, updated_at="2026-01-01T00:00:00Z", recall_at="", workflow_restart_required=0, receipts=[dict(recipient_platform="tmall", release_no=1, received_at=stamp)])
        item.update(updates)
        return item

    def sync(self, *items):
        rot.sync_sources(self.path, list(items), "tester")

    def rows(self, **criteria):
        return rot.candidates(self.path, criteria, date(2026, 10, 20))

    def new_plan(self, criteria=None, target="唯品"):
        criteria = criteria or {"mode": "time", "min_days": 70}
        return rot.create_plan(self.path, criteria, None, target, "tester", rot.fingerprint(rot.candidates(self.path, criteria)))

    def baseline(self, path, exclude_rotation=True):
        with sqlite3.connect(path) as connection:
            tables = [r[0] for r in connection.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'") if not exclude_rotation or not r[0].startswith("rotation_")]
            return {name: connection.execute(f'SELECT * FROM "{name}" ORDER BY rowid').fetchall() for name in tables}

    def test_only_snapshot_tmall_vip_and_earliest_style_color_and_system(self):
        one = self.source(channel="同款", receipts=[
            dict(recipient_platform="all", release_no=1, received_at="2025-01-01T00:00:00Z"),
            dict(recipient_platform="design", release_no=1, received_at="2025-01-02T00:00:00Z"),
            dict(recipient_platform="vip", release_no=1, received_at="2026-03-01T00:00:00Z"),
        ])
        two = self.source(2, channel="同款", stamp="2026-02-01T00:00:00Z")
        self.sync(one, two)
        row = self.rows()[0]
        self.assertEqual(row["reference_at"], "2026-02-01T00:00:00Z")
        self.assertEqual(row["source_ids"], [1, 2])
        self.assertEqual(len(row["receipt_evidence"]), 2)
        self.sync(self.source(receipts=[dict(recipient_platform="all", release_no=1, received_at="2025-01-01T00:00:00Z")]))
        self.assertEqual(self.rows(), [])

    def test_latest_recall_of_any_color_resets_entire_style(self):
        old_color = self.source(1)
        recalled = self.source(2, c_release_no=2, recall_at="2026-09-01T00:00:00Z", receipts=[dict(recipient_platform="tmall", release_no=1, received_at="2026-01-01T00:00:00Z")])
        self.sync(old_color, recalled)
        self.assertEqual(self.rows(min_days=0), [])
        recalled["receipts"].append(dict(recipient_platform="vip", release_no=2, received_at="2026-09-05T00:00:00Z"))
        self.sync(old_color, recalled)
        self.assertEqual(self.rows(min_days=0)[0]["reference_at"], "2026-09-05T00:00:00Z")
        self.assertEqual(self.rows(min_days=70), [])

    def test_shanghai_natural_days_and_receipt_day_zero(self):
        self.sync(self.source(stamp="2026-08-10T16:01:00Z"))
        self.assertEqual(self.rows(min_days=70)[0]["age_days"], 70)
        self.assertEqual(self.rows(min_days=71), [])
        rows = rot.candidates(self.path, {"min_days": 0}, date(2026, 8, 11))
        self.assertEqual(rows[0]["age_days"], 0)

    def test_full_style_excludes_recalled_archived_or_mixed_channel(self):
        for updates in ({"status": "pending"}, {"lifecycle_status": "archived"}, {"launch_channel": "唯品"}, {"workflow_restart_required": 1}):
            self.sync(self.source(), self.source(2, **updates))
            self.assertEqual(self.rows(), [], updates)

    def test_targets_exclude_intra_tmall_and_single_color(self):
        for c in rot.TMALL:
            self.assertTrue(rot.valid_target(c, "唯品"))
            for target in rot.TMALL:
                self.assertFalse(rot.valid_target(c, target))
        self.assertTrue(rot.valid_target("同款", "天猫奥"))
        self.assertFalse(rot.valid_target("同款", "同款"))
        self.sync(self.source(), self.source(2))
        with self.assertRaises(ValueError):
            self.new_plan(target="天猫奥")
        pid = self.new_plan()
        self.assertEqual(rot.plan_detail(self.path, pid)["items"][0]["snapshot"]["source_ids"], [1, 2])

    def sales_file(self, rows):
        return rot.workbook_bytes(rot.SALES_HEADERS, rows, "轮转销量")

    def test_sales_missing_not_zero_and_same_period_for_both_systems(self):
        self.sync(self.source(channel="同款"))
        criteria = dict(mode="sales", start_date="2026-07-01", end_date="2026-07-31", max_sales=0)
        self.assertEqual(self.rows(**criteria), [])
        rot.import_sales(self.path, self.sales_file([["A001", "天猫体系", "2026-07-01", "2026-07-31", 0]]), "tester")
        self.assertEqual(self.rows(**criteria), [])
        rot.import_sales(self.path, self.sales_file([["A001", "唯品体系", "2026-07-01", "2026-07-31", 0]]), "tester")
        row = self.rows(**criteria)[0]
        self.assertEqual(row["sales_quantity"], 0)
        self.assertEqual(len(row["sales"]), 2)
        self.assertEqual(self.rows(**{**criteria, "end_date": "2026-07-30"}), [])

    def test_sales_duplicates_overlaps_replace_and_atomic_validation(self):
        self.sync(self.source())
        content = self.sales_file([["A001", "天猫体系", "2026-07-01", "2026-07-31", 2]])
        rot.import_sales(self.path, content, "tester")
        with self.assertRaisesRegex(ValueError, "已存在或重叠"):
            rot.import_sales(self.path, content, "tester")
        with self.assertRaises(ValueError):
            rot.import_sales(self.path, self.sales_file([["A001", "天猫体系", "2026-07-02", "2026-07-31", 2]]), "tester", True)
        rot.import_sales(self.path, content, "tester", True)
        with db.get_connection(self.path) as connection:
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM rotation_sales").fetchone()[0], 2)
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM rotation_sales WHERE active=1").fetchone()[0], 1)
        for bad in (None, 1.2, "=1+2", True):
            with self.assertRaises(ValueError):
                rot.import_sales(self.path, self.sales_file([["A001", "唯品体系", "2026-07-01", "2026-07-31", bad]]), "tester")
        with self.assertRaisesRegex(ValueError, "重复"):
            rot.import_sales(self.path, self.sales_file([["A001", "唯品体系", "2026-07-01", "2026-07-31", 1]] * 2), "tester")
        with self.assertRaises(ValueError):
            rot.import_sales(self.path, self.sales_file([["A001", "天猫体系", "2026-07-01", "2026-07-31", 5], ["UNKNOWN", "唯品体系", "2026-07-01", "2026-07-31", 1]]), "tester", True)
        with db.get_connection(self.path) as connection:
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM rotation_sales WHERE active=1").fetchone()[0], 1)

    def test_time_sales_and_multiselect_filters(self):
        self.sync(self.source(), self.source(2, code="B001", supplier="供应商B", category="连衣裙"))
        self.assertEqual(len(self.rows(supplier=["供应商A", "供应商B"])), 2)
        self.assertEqual(len(self.rows(category=["连衣裙"])), 1)
        rot.import_sales(self.path, self.sales_file([["A001", "天猫体系", "2026-07-01", "2026-07-31", 5]]), "tester")
        args = dict(mode="both", start_date="2026-07-01", end_date="2026-07-31", max_sales=5)
        self.assertEqual(len(self.rows(**args)), 1)
        self.assertEqual(self.rows(**{**args, "max_sales": 4}), [])
        self.assertEqual(self.rows(**{**args, "min_days": 999}), [])

    def test_plan_lifecycle_permissions_duplicate_snapshot_and_data_preservation(self):
        before = self.baseline(self.path)
        self.sync(self.source(), self.source(2))
        pid = self.new_plan()
        item = rot.plan_detail(self.path, pid)["items"][0]
        self.assertEqual(self.rows(), [])
        with self.assertRaises(PermissionError):
            rot.update_plan(self.path, pid, "confirm", self.planner, {})
        rot.update_plan(self.path, pid, "save", self.planner, {f"target_{item['id']}": "唯品"})
        rot.update_plan(self.path, pid, "confirm", self.admin, {})
        rot.update_plan(self.path, pid, "start", self.planner, {})
        rot.update_plan(self.path, pid, "feedback", self.planner, {f"done_{item['id']}": "1", "feedback": "人工执行测试"})
        plan = rot.plan_detail(self.path, pid)
        self.assertEqual(plan["state"], "completed")
        self.assertGreaterEqual(len(plan["events"]), 5)
        self.assertEqual(plan["items"][0]["snapshot"]["colors"], ["A001-1", "A001-2"])
        self.assertEqual(before, self.baseline(self.path))

    def test_stale_draft_and_failed_sync_do_not_overwrite_history(self):
        self.sync(self.source())
        pid = self.new_plan()
        self.sync(self.source(channel="唯品"))
        with self.assertRaisesRegex(ValueError, "来源资料已变化"):
            rot.update_plan(self.path, pid, "confirm", self.admin, {})
        self.assertEqual(rot.plan_detail(self.path, pid)["items"][0]["snapshot"]["channel"], "天猫官")
        with patch("planning_center.rotation_web.urlopen", side_effect=OSError("offline")):
            with self.assertRaises(ValueError):
                self.app.sync_rotation(self.planner)
        self.assertEqual(self.rows(), [])  # Still guarded by the open draft.
        rot.update_plan(self.path, pid, "cancel", self.planner, {})
        self.assertEqual(self.rows()[0]["channel"], "唯品")

    def test_cross_page_selection_1000_styles_and_no_n_plus_one(self):
        self.sync(*(self.source(i, code=f"A{i:05}") for i in range(1, 1001)))
        started = time.monotonic()
        pid = self.new_plan()
        self.assertEqual(len(rot.plan_detail(self.path, pid)["items"]), 1000)
        self.assertLess(time.monotonic() - started, 6)

    def test_workbooks_unprotected_formula_safe_and_wps_editable(self):
        content = rot.workbook_bytes(("款号", "数量"), [("=HYPERLINK(1)", 1)], "执行清单")
        book = load_workbook(io.BytesIO(content))
        self.assertFalse(book.active.protection.sheet)
        self.assertEqual(book.active["A2"].data_type, "s")
        self.assertFalse(book.active["A2"].protection.locked)
        book.active["A2"] = "可编辑"
        book.save(io.BytesIO())

    def test_wsgi_routes_csrf_role_and_export(self):
        cookie = self.login_cookie(self.app, "planner")
        self.assertEqual(self.wsgi_request(self.app, "/rotation")["status"], "302 Found")
        self.assertTrue(self.wsgi_request(self.app, "/rotation", cookie=cookie)["status"].startswith("200"))
        bad = self.wsgi_request(self.app, "/rotation/sync", method="POST", cookie=cookie)
        self.assertTrue(bad["status"].startswith("403"))
        self.sync(self.source())
        rows = rot.candidates(self.path, {"min_days": 70})
        form = dict(csrf=self.app.rotation_csrf(self.planner), criteria=json.dumps({"min_days": 70}), scope="filtered", target="唯品", signature=rot.fingerprint(rows))
        response = self.wsgi_request(self.app, "/rotation/create", method="POST", cookie=cookie, body=urlencode(form).encode())
        self.assertEqual(dict(response["headers"])["Location"], "/rotation/1")
        self.assertTrue(self.wsgi_request(self.app, "/rotation/1", cookie=cookie)["status"].startswith("200"))
        export = self.wsgi_request(self.app, "/rotation/1/export.xlsx", cookie=cookie)
        self.assertFalse(load_workbook(io.BytesIO(export["body"])).active.protection.sheet)
        self.assertIn("货品轮转规则", self.wsgi_request(self.app, "/system-rules", cookie=cookie)["body"].decode())

    def test_existing_db_backup_is_verified_and_migration_idempotent(self):
        path = self.root / "legacy.db"
        with sqlite3.connect(path) as connection:
            connection.execute("CREATE TABLE sentinel(id INTEGER PRIMARY KEY,value TEXT)")
            connection.execute("INSERT INTO sentinel VALUES(1,'existing data')")
        rot.ensure_schema(path)
        backups = list((self.root / "backups").glob("pre-rotation-*.db"))
        self.assertEqual(len(backups), 1)
        with sqlite3.connect(backups[0]) as connection:
            self.assertEqual(connection.execute("SELECT value FROM sentinel").fetchone()[0], "existing data")
            self.assertEqual(connection.execute("PRAGMA quick_check").fetchone()[0], "ok")
        rot.ensure_schema(path)
        self.assertEqual(len(list((self.root / "backups").glob("pre-rotation-*.db"))), 1)
        with sqlite3.connect(path) as connection:
            self.assertEqual(connection.execute("SELECT value FROM sentinel").fetchone()[0], "existing data")

    def test_backup_failure_stops_additive_migration(self):
        path = self.root / "legacy-fail.db"
        with sqlite3.connect(path) as connection:
            connection.execute("CREATE TABLE sentinel(id INTEGER PRIMARY KEY)")
            connection.execute("INSERT INTO sentinel VALUES(1)")
        with patch("planning_center.rotation.Path.touch", side_effect=OSError("backup unavailable")):
            with self.assertRaises(OSError):
                rot.ensure_schema(path)
        with sqlite3.connect(path) as connection:
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM sentinel").fetchone()[0], 1)
            self.assertIsNone(connection.execute("SELECT 1 FROM sqlite_master WHERE name='rotation_sources'").fetchone())

    def test_identity_rename_cannot_create_second_active_plan_or_reuse_sales(self):
        self.sync(self.source())
        rot.import_sales(self.path, self.sales_file([["A001", "天猫体系", "2026-07-01", "2026-07-31", 0]]), "tester")
        pid = self.new_plan()
        self.sync(self.source(code="RENAMED"))
        self.assertEqual(self.rows(), [])
        rot.update_plan(self.path, pid, "cancel", self.planner, {})
        self.assertEqual(self.rows()[0]["style_code"], "RENAMED")
        self.assertEqual(self.rows(mode="sales", start_date="2026-07-01", end_date="2026-07-31", max_sales=0), [])
        self.sync(self.source(), self.source(2))
        self.assertEqual(self.rows(mode="sales", start_date="2026-07-01", end_date="2026-07-31", max_sales=0), [])

    def test_concurrent_create_has_only_one_winner(self):
        self.sync(self.source(), self.source(2))
        criteria = {"min_days": 70}
        signature = rot.fingerprint(rot.candidates(self.path, criteria))
        def create(_):
            try:
                return rot.create_plan(self.path, criteria, None, "唯品", "tester", signature)
            except ValueError:
                return None
        with ThreadPoolExecutor(max_workers=4) as pool:
            ids = list(pool.map(create, range(4)))
        self.assertEqual(sum(i is not None for i in ids), 1)

    def test_changed_sales_or_unsaved_target_cannot_confirm(self):
        self.sync(self.source(channel="唯品"))
        content = self.sales_file([["A001", "唯品体系", "2026-07-01", "2026-07-31", 1]])
        rot.import_sales(self.path, content, "tester")
        pid = self.new_plan(dict(mode="sales", start_date="2026-07-01", end_date="2026-07-31", max_sales=1), "天猫官")
        item = rot.plan_detail(self.path, pid)["items"][0]
        with self.assertRaisesRegex(ValueError, "未保存"):
            rot.update_plan(self.path, pid, "confirm", self.admin, {f"target_{item['id']}": "天猫奥"})
        rot.import_sales(self.path, content, "tester", True)
        with self.assertRaisesRegex(ValueError, "销售资料已"):
            rot.update_plan(self.path, pid, "confirm", self.admin, {})

    def test_paginated_sync_rejects_mixed_revision_without_mutation(self):
        self.sync(self.source())
        before = self.baseline(self.path, False)
        responses = [dict(schema_version=1, revision="one", items=[self.source(2)], has_more=True, next_after=2), dict(schema_version=1, revision="two", items=[], has_more=False, next_after=2)]
        def opened(*args, **kwargs):
            return closing(io.BytesIO(json.dumps(responses.pop(0)).encode()))
        with patch("planning_center.rotation_web.urlopen", side_effect=opened):
            with self.assertRaisesRegex(ValueError, "同步期间"):
                self.app.sync_rotation(self.planner)
        self.assertEqual(before, self.baseline(self.path, False))

    def test_executed_style_waits_for_new_channel_or_release_round(self):
        self.sync(self.source(), self.source(2))
        pid = self.new_plan()
        item = rot.plan_detail(self.path, pid)["items"][0]
        rot.update_plan(self.path, pid, "confirm", self.admin, {})
        rot.update_plan(self.path, pid, "feedback", self.planner, {f"done_{item['id']}": "1", "feedback": "已人工执行整款"})
        self.assertEqual(self.rows(), [])
        self.sync(self.source(product_name="新名称"), self.source(2))
        self.assertEqual(self.rows(), [])
        self.sync(self.source(channel="唯品"), self.source(2, channel="唯品"))
        self.assertEqual(len(self.rows()), 1)
        self.sync(*(self.source(i, c_release_no=2, recall_at="2026-05-01T00:00:00Z", receipts=[dict(recipient_platform="tmall", release_no=2, received_at="2026-06-01T00:00:00Z")]) for i in (1, 2)))
        self.assertEqual(len(self.rows()), 1)
        self.assertEqual(rot.plan_detail(self.path, pid)["state"], "completed")

    def test_stale_plan_form_cannot_overwrite_another_edit(self):
        self.sync(self.source(channel="唯品"))
        pid = self.new_plan(target="天猫官")
        plan = rot.plan_detail(self.path, pid)
        token = rot.plan_signature(plan)
        iid = plan["items"][0]["id"]
        rot.update_plan(self.path, pid, "save", self.planner, {"plan_signature": token, f"target_{iid}": "天猫奥"})
        before = self.baseline(self.path, False)
        with self.assertRaisesRegex(ValueError, "其他同事更新"):
            rot.update_plan(self.path, pid, "confirm", self.admin, {"plan_signature": token})
        self.assertEqual(before, self.baseline(self.path, False))
        plan = rot.plan_detail(self.path, pid)
        self.assertEqual(plan["items"][0]["target"], "天猫奥")
        rot.update_plan(self.path, pid, "confirm", self.admin, {"plan_signature": rot.plan_signature(plan)})
        self.assertEqual(rot.plan_detail(self.path, pid)["state"], "confirmed")

    def test_readiness_requires_existing_upgrade_backup_evidence(self):
        path = self.root / "readiness.db"
        with sqlite3.connect(path) as connection:
            connection.execute("CREATE TABLE sentinel(id INTEGER PRIMARY KEY)")
        rot.ensure_schema(path)
        self.assertEqual(rot.readiness(path), {"schema_ready": True, "migration_backup_verified": True, "migration_backup_required": True})
        with patch("planning_center.rotation.Path.is_file", return_value=False):
            self.assertFalse(rot.readiness(path)["migration_backup_verified"])
        self.assertTrue(rot.readiness(path)["migration_backup_verified"])

    def test_invalid_sync_is_atomic_and_preserves_all_existing_tables(self):
        self.sync(self.source())
        before = self.baseline(self.path, False)
        for invalid in (self.source(2, receipts=None), self.source(2, recall_at="bad-date"), self.source(2, receipts=[dict(release_no=1, received_at="")])):
            with self.assertRaises(ValueError):
                rot.sync_sources(self.path, [self.source(3), invalid], "tester")
            self.assertEqual(before, self.baseline(self.path, False))

    def test_catalog_feed_is_token_protected_read_only_and_detects_unreceived_recall(self):
        path = self.root / "catalog.db"
        catalog_db.init_db(path)
        app = CatalogApplication(path, self.root / "uploads", planning_api_token="test-token")
        before = self.baseline(path, False)
        route = "/api/internal/planning/rotation-products?limit=1"
        self.assertTrue(self.wsgi_request(app, route)["status"].startswith("401"))
        response = self.wsgi_request(app, route, authorization="Bearer test-token")
        payload = json.loads(response["body"])
        self.assertEqual(len(payload["items"]), 1)
        self.assertTrue(payload["has_more"])
        self.assertEqual(before, self.baseline(path, False))
        pid = payload["items"][0]["id"]
        with catalog_db.get_connection(path) as connection:
            uid = connection.execute("SELECT id FROM users LIMIT 1").fetchone()[0]
            for status, stamp in (("published", "2026-07-01T00:00:00Z"), ("pending", "2026-07-02T00:00:00Z")):
                connection.execute("INSERT INTO product_logs(product_id,actor_user_id,action,action_label,details,created_at) VALUES(?,?,?,?,?,?)", (pid, uid, "status:" + status, "test", "test", stamp))
        self.assertEqual(rot.catalog_page(path)["items"][0]["recall_at"], "2026-07-02T00:00:00Z")

    def test_real_catalog_receive_recall_republish_resets_whole_style_clock(self):
        path = self.root / "catalog-loop.db"
        catalog_db.init_db(path, seed_samples=False)
        with catalog_db.get_connection(path) as connection:
            users = {r["username"]: dict(r) for r in connection.execute("SELECT * FROM users")}
            ids = []
            with patch("catalog_backend.db.utc_now", return_value="2026-01-01T01:00:00Z"):
                for color in ("黑", "白"):
                    pid = catalog_db.create_product(connection, dict(style_code="LOOP01", style_color="LOOP01-" + color, product_name="针织衫", launch_channel="同款", launch_price="599", tax_included_price="150"), users["a_editor"]["id"], "A")
                    catalog_db.change_product_status(connection, pid, "pending", users["a_editor"]["id"], "协作", "测试")
                    catalog_db.change_product_status(connection, pid, "published", users["b_editor"]["id"], "提交运营", "测试")
                    self.assertTrue(catalog_db.record_c_product_receipt(connection, pid, users["c_viewer"]["id"], 1))
                    ids.append(pid)
        self.sync(*rot.catalog_page(path)["items"])
        self.assertEqual(self.rows()[0]["source_ids"], ids)
        with catalog_db.get_connection(path) as connection:
            with patch("catalog_backend.db.utc_now", return_value="2026-09-01T01:00:00Z"):
                catalog_db.change_product_status(connection, ids[1], "pending", users["a_editor"]["id"], "召回", "测试")
        self.sync(*rot.catalog_page(path)["items"])
        self.assertEqual(self.rows(min_days=0), [])
        with catalog_db.get_connection(path) as connection:
            with patch("catalog_backend.db.utc_now", return_value="2026-09-02T01:00:00Z"):
                catalog_db.change_product_status(connection, ids[1], "published", users["b_editor"]["id"], "重新提交运营", "测试")
        self.sync(*rot.catalog_page(path)["items"])
        self.assertEqual(self.rows(min_days=0), [])
        with catalog_db.get_connection(path) as connection:
            with patch("catalog_backend.db.utc_now", return_value="2026-09-05T01:00:00Z"):
                self.assertTrue(catalog_db.record_c_product_receipt(connection, ids[1], users["c_viewer"]["id"], 2))
        before = self.baseline(path, False)
        self.sync(*rot.catalog_page(path)["items"])
        row = self.rows(min_days=0)[0]
        self.assertEqual(row["reference_at"], "2026-09-05T01:00:00Z")
        self.assertEqual(row["age_days"], 45)
        self.assertEqual(self.rows(min_days=70), [])
        self.assertEqual(before, self.baseline(path, False))


if __name__ == "__main__":
    unittest.main()
